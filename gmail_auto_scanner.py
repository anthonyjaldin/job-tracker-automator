#!/usr/bin/env python3
"""
Real Gmail Scanner - Uses App Password
"""
import imaplib
import email
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

EMAIL = "a.jaldin18@gmail.com"
APP_PASSWORD = "luxuxzabsziytpxt"

COMPANIES = {
    'amazon': 'Amazon',
    'microsoft': 'Microsoft',
    'google': 'Google',
    'capital one': 'Capital One',
    'jpmorgan': 'JPMorgan',
    'apple': 'Apple',
    'meta': 'Meta',
    'tesla': 'Tesla',
}

def connect_gmail():
    """Connect to Gmail"""
    print("🔐 Connecting to Gmail...")
    try:
        imap = imaplib.IMAP4_SSL("imap.gmail.com")
        imap.login(EMAIL, APP_PASSWORD)
        print("✅ Connected!")
        return imap
    except Exception as e:
        print(f"❌ Error: {e}")
        return None

def scan_emails(imap):
    """Scan sent emails"""
    print("\n📧 Scanning sent emails...")
    
    try:
        # Use correct Gmail label
        imap.select('"[Gmail]/Sent Mail"')
        status, messages = imap.search(None, 'ALL')
        
        msg_ids = messages[0].split()
        applications = []
        
        for msg_id in msg_ids[-50:]:
            try:
                status, msg_data = imap.fetch(msg_id, '(RFC822)')
                msg = email.message_from_bytes(msg_data[0][1])
                
                subject = msg.get('subject', '').lower()
                date_str = msg.get('date', '')
                
                for key, company in COMPANIES.items():
                    if key in subject:
                        date_display = date_str[:10] if date_str else datetime.now().strftime('%m/%d/%Y')
                        applications.append({
                            'date': date_display,
                            'company': company,
                            'subject': subject[:50],
                            'status': 'Applied'
                        })
                        print(f"  ✓ {company}")
                        break
            except:
                pass
        
        return applications
    
    except Exception as e:
        print(f"❌ Error: {e}")
        return []

def update_excel(applications):
    """Update Excel"""
    print(f"\n📊 Adding {len(applications)} to Excel...")
    
    try:
        wb = load_workbook('job_applications_tracker.xlsx')
        ws = wb.active
        
        row = 2
        for app in applications:
            ws[f'A{row}'] = app['date']
            ws[f'B{row}'] = app['company']
            ws[f'C{row}'] = app['subject']
            ws[f'D{row}'] = app['status']
            ws[f'E{row}'] = datetime.now().strftime('%m/%d/%Y')
            
            ws[f'D{row}'].fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type="solid")
            row += 1
        
        wb.save('job_applications_tracker.xlsx')
        print("✅ Excel updated!")
    except Exception as e:
        print(f"❌ Error: {e}")

def main():
    print("🚀 Gmail Auto Scanner")
    print("=" * 60)
    
    imap = connect_gmail()
    if not imap:
        return
    
    apps = scan_emails(imap)
    
    try:
        imap.close()
        imap.logout()
    except:
        pass
    
    if apps:
        update_excel(apps)
    else:
        print("⚠️  No job applications found")
    
    print("=" * 60)
    print("✅ Done!")

if __name__ == '__main__':
    main()
