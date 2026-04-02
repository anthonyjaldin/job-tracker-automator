#!/usr/bin/env python3
"""
Simple Job Application Scanner - Run manually whenever you want
"""
import os
import json
from datetime import datetime

try:
    import base64
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
except ImportError:
    print("Installing packages...")
    os.system("pip install google-auth-oauthlib google-api-python-client openpyxl --break-system-packages")
    import base64
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']

# Companies to look for
COMPANIES = {
    'amazon': 'Amazon',
    'microsoft': 'Microsoft',
    'google': 'Google',
    'apple': 'Apple',
    'meta': 'Meta',
    'tesla': 'Tesla',
    'ibm': 'IBM',
    'oracle': 'Oracle',
    'cisco': 'Cisco',
    'intel': 'Intel',
    'capital one': 'Capital One',
    'jpmorgan': 'JPMorgan',
    'goldman': 'Goldman Sachs',
    'deloitte': 'Deloitte',
    'accenture': 'Accenture',
    'aws': 'Amazon AWS',
    'azure': 'Microsoft Azure'
}

def get_gmail_service():
    """Authenticate and get Gmail service"""
    creds = None
    
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
    
    if not creds or not creds.valid:
        flow = InstalledAppFlow.from_client_secrets_file('credentials.json', SCOPES)
        creds = flow.run_local_server(port=0)
        with open('token.json', 'w') as token:
            token.write(creds.to_json())
    
    return build('gmail', 'v1', credentials=creds)

def scan_emails(service):
    """Scan sent emails for job applications"""
    print("📧 Scanning your sent emails...")
    
    try:
        results = service.users().messages().list(
            userId='me',
            q='in:sent',
            maxResults=100
        ).execute()
        
        messages = results.get('messages', [])
        applications = []
        
        for msg in messages:
            msg_data = service.users().messages().get(userId='me', id=msg['id']).execute()
            headers = msg_data['payload'].get('headers', [])
            
            subject = next((h['value'] for h in headers if h['name'] == 'Subject'), '').lower()
            date_str = next((h['value'] for h in headers if h['name'] == 'Date'), '')
            
            # Extract date (rough format)
            date_parts = date_str.split(',')
            if len(date_parts) > 1:
                date_clean = date_parts[1].strip().split()[0:3]
                date_display = ' '.join(date_clean) if date_clean else date_str[:10]
            else:
                date_display = date_str[:10]
            
            # Check for company mentions
            for company_key, company_name in COMPANIES.items():
                if company_key in subject:
                    applications.append({
                        'date': date_display,
                        'company': company_name,
                        'subject': subject.title(),
                        'status': 'Applied'
                    })
                    print(f"  ✓ Found: {company_name}")
                    break
        
        return applications
    
    except Exception as e:
        print(f"❌ Error scanning emails: {e}")
        return []

def update_excel(applications):
    """Update Excel with found applications"""
    if not os.path.exists('job_applications_tracker.xlsx'):
        print("❌ Excel file not found!")
        return
    
    print(f"\n📊 Updating Excel with {len(applications)} applications...")
    
    wb = load_workbook('job_applications_tracker.xlsx')
    ws = wb.active
    
    # Status colors
    status_colors = {
        'Applied': 'D4EDDA',
        'Pending': 'FFF3CD',
        'Interview': 'B8DAFF',
        'Rejected': 'F8D7DA',
        'Offer': 'D1E7DD'
    }
    
    # Add applications starting at row 2
    row_num = 2
    for app in applications:
        ws[f'A{row_num}'] = app['date']
        ws[f'B{row_num}'] = app['company']
        ws[f'C{row_num}'] = app['subject']
        ws[f'D{row_num}'] = app['status']
        ws[f'E{row_num}'] = datetime.now().strftime('%m/%d/%Y')
        
        # Color code
        status_fill = PatternFill(
            start_color=status_colors.get(app['status'], 'FFFFFF'),
            end_color=status_colors.get(app['status'], 'FFFFFF'),
            fill_type="solid"
        )
        ws[f'D{row_num}'].fill = status_fill
        
        row_num += 1
    
    wb.save('job_applications_tracker.xlsx')
    print(f"✅ Excel updated with {len(applications)} applications!")

def main():
    print("🚀 Job Application Scanner")
    print("=" * 60)
    
    if not os.path.exists('credentials.json'):
        print("❌ credentials.json not found!")
        return
    
    try:
        print("🔐 Authenticating with Gmail...")
        service = get_gmail_service()
        print("✅ Gmail authenticated!")
        
        applications = scan_emails(service)
        
        if applications:
            update_excel(applications)
            print("=" * 60)
            print("✅ Done! Check: job_applications_tracker.xlsx")
        else:
            print("⚠️  No job applications found in sent emails")
        
    except Exception as e:
        print(f"❌ Error: {e}")

if __name__ == '__main__':
    main()
