#!/usr/bin/env python3
"""
Job Application Scanner - Fixed OAuth
"""
import os
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from datetime import datetime

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']

COMPANIES = {
    'amazon': 'Amazon',
    'microsoft': 'Microsoft',
    'google': 'Google',
    'capital one': 'Capital One',
    'jpmorgan': 'JPMorgan',
    'apple': 'Apple',
}

def authenticate():
    """Authenticate with Gmail"""
    creds = None
    
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
        if creds.valid:
            return creds
    
    flow = InstalledAppFlow.from_client_secrets_file(
        'credentials.json',
        SCOPES
    )
    
    creds = flow.run_local_server(port=8080, open_browser=True)
    
    with open('token.json', 'w') as token:
        token.write(creds.to_json())
    
    return creds

def scan_emails(creds):
    """Scan Gmail"""
    print("\n📧 Scanning emails...")
    service = build('gmail', 'v1', credentials=creds)
    
    results = service.users().messages().list(userId='me', q='in:sent', maxResults=50).execute()
    messages = results.get('messages', [])
    
    applications = []
    for msg in messages:
        msg_data = service.users().messages().get(userId='me', id=msg['id']).execute()
        headers = msg_data['payload'].get('headers', [])
        
        subject = next((h['value'] for h in headers if h['name'] == 'Subject'), '').lower()
        date_str = next((h['value'] for h in headers if h['name'] == 'Date'), '')
        
        for key, company in COMPANIES.items():
            if key in subject:
                applications.append({
                    'date': date_str[:10],
                    'company': company,
                    'subject': subject[:50],
                    'status': 'Applied'
                })
                print(f"  ✓ {company}")
                break
    
    return applications

def update_excel(applications):
    """Update Excel"""
    wb = load_workbook('job_applications_tracker.xlsx')
    ws = wb.active
    
    row = 2
    for app in applications:
        ws[f'A{row}'] = app['date']
        ws[f'B{row}'] = app['company']
        ws[f'C{row}'] = app['subject']
        ws[f'D{row}'] = app['status']
        ws[f'E{row}'] = datetime.now().strftime('%m/%d/%Y')
        
        fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type="solid")
        ws[f'D{row}'].fill = fill
        row += 1
    
    wb.save('job_applications_tracker.xlsx')
    print(f"✅ Updated Excel with {len(applications)} apps!")

def main():
    print("🚀 Job Scanner")
    print("=" * 60)
    
    try:
        print("🔐 Authenticating...")
        creds = authenticate()
        print("✅ Authenticated!")
        
        apps = scan_emails(creds)
        if apps:
            update_excel(apps)
        
        print("=" * 60)
        print("✅ Done!")
    except Exception as e:
        print(f"❌ {e}")

if __name__ == '__main__':
    main()
